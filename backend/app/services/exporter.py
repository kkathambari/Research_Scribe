import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.config import settings
from app.services.citation_formatter import citation_formatter

class ManuscriptExporterService:
    """
    Subsystem for template-preserving document assembly and multi-format export
    conforming to SRS v3.0 (FR-23, FR-26).
    """

    def export_manuscript(
        self,
        manuscript: Dict[str, Any],
        sections: List[Dict[str, Any]],
        sources: List[Dict[str, Any]],
        claims: Optional[List[Dict[str, Any]]] = None,
        export_format: str = "docx",
        citation_style: str = "IEEE",
        include_verification_appendix: bool = False,
        template_path: Optional[str] = None
    ) -> Path:
        """
        Exports manuscript into DOCX (template-preserved), Markdown, or BibTeX.
        """
        export_format = export_format.lower()
        title = manuscript.get("title") or "Research Manuscript"
        clean_name = "".join(c for c in title if c.isalnum() or c in (' ', '_', '-')).rstrip()
        clean_name = clean_name.replace(' ', '_')[:40] or "manuscript"
        manuscript_id = manuscript.get("id", "draft")

        if export_format == "docx":
            filename = f"{clean_name}_{manuscript_id[:8]}.docx"
            out_path = settings.EXPORTS_DIR / filename
            self._generate_docx(
                out_path,
                title,
                sections,
                sources,
                claims,
                citation_style,
                include_verification_appendix,
                template_path,
                abstract=manuscript.get("abstract", ""),
                target_journal=manuscript.get("target_journal", "")
            )
            return out_path

        elif export_format == "pdf":
            filename = f"{clean_name}_{manuscript_id[:8]}.pdf"
            out_path = settings.EXPORTS_DIR / filename
            self._generate_pdf(
                out_path,
                title,
                sections,
                sources,
                claims,
                citation_style,
                include_verification_appendix,
                abstract=manuscript.get("abstract", ""),
                target_journal=manuscript.get("target_journal", "")
            )
            return out_path

        elif export_format == "bibtex":
            filename = f"{clean_name}_{manuscript_id[:8]}.bib"
            out_path = settings.EXPORTS_DIR / filename
            bibtex_content = citation_formatter.generate_bibtex(sources)
            out_path.write_text(bibtex_content, encoding="utf-8")
            return out_path

        else:  # markdown default
            filename = f"{clean_name}_{manuscript_id[:8]}.md"
            out_path = settings.EXPORTS_DIR / filename
            md_content = self._generate_markdown(title, sections, sources, claims, citation_style, include_verification_appendix)
            out_path.write_text(md_content, encoding="utf-8")
            return out_path

    def _generate_docx(
        self,
        out_path: Path,
        title: str,
        sections: List[Dict[str, Any]],
        sources: List[Dict[str, Any]],
        claims: Optional[List[Dict[str, Any]]],
        citation_style: str,
        include_verification_appendix: bool,
        template_path: Optional[str] = None,
        abstract: str = "",
        target_journal: str = ""
    ):
        """FR-23: Template-preserving DOCX assembly retaining fonts, margins, headers, and footers."""
        is_template_loaded = False
        if template_path and Path(template_path).exists():
            try:
                doc = docx.Document(template_path)
                is_template_loaded = True
            except Exception:
                doc = docx.Document()
        else:
            doc = docx.Document()

        if is_template_loaded:
            # Map of placeholders to section contents
            sec_by_type = {s.get("section_type", "").lower(): s for s in sections}
            replaced_types = set()

            def _replace_in_paragraph_preserving_runs(p, pattern: str, replacement: str) -> bool:
                """Replaces pattern in paragraph while preserving individual run styles and formatting."""
                if not p.text or not re.search(pattern, p.text, flags=re.IGNORECASE):
                    return False
                # If pattern exists inside a single run, replace strictly within that run to retain formatting
                for r in p.runs:
                    if re.search(pattern, r.text, flags=re.IGNORECASE):
                        r.text = re.sub(pattern, replacement, r.text, flags=re.IGNORECASE)
                        return True
                # If placeholder spans multiple runs, replace text while retaining first run's font/style
                full_text = p.text
                new_full = re.sub(pattern, replacement, full_text, flags=re.IGNORECASE)
                if p.runs:
                    p.runs[0].text = new_full
                    for r in p.runs[1:]:
                        r.text = ""
                else:
                    p.text = new_full
                return True

            def _sub_placeholders(p):
                if not p.text:
                    return

                _replace_in_paragraph_preserving_runs(p, r'\{\{TITLE\}\}|\[paper title\]|\[title of paper\]', title)

                abs_text = sec_by_type.get("abstract", {}).get("content") or abstract or "This manuscript presents an evidence-grounded investigation."
                if _replace_in_paragraph_preserving_runs(p, r'\{\{ABSTRACT\}\}|\[abstract\]', abs_text):
                    replaced_types.add("abstract")

                _replace_in_paragraph_preserving_runs(p, r'\{\{AUTHORS\}\}|\[author\]|\[authors\]|\[author names\]', "Research Author(s)")
                _replace_in_paragraph_preserving_runs(p, r'\{\{TARGET_JOURNAL\}\}|\{\{JOURNAL\}\}|\[target publication\]|\[journal\]', target_journal or "IEEE Transactions")

                intro_sec = sec_by_type.get("introduction")
                if intro_sec and _replace_in_paragraph_preserving_runs(p, r'\{\{INTRODUCTION\}\}|\[introduction\]', intro_sec.get("content", "")):
                    replaced_types.add("introduction")

                meth_sec = sec_by_type.get("methodology")
                if meth_sec and _replace_in_paragraph_preserving_runs(p, r'\{\{METHODOLOGY\}\}|\[methodology\]', meth_sec.get("content", "")):
                    replaced_types.add("methodology")

                _replace_in_paragraph_preserving_runs(p, r'\{\{KEYWORDS\}\}|\[keywords\]', "Keywords: research automation; repository analysis; evidence provenance; reproducible science")

            # 1. Main body paragraphs
            for p in doc.paragraphs:
                _sub_placeholders(p)

            # 2. Table cells (for journal templates with metadata / header tables)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            _sub_placeholders(p)

            # 3. Headers and footers across sections
            for s in doc.sections:
                if s.header:
                    for p in s.header.paragraphs:
                        _sub_placeholders(p)
                if s.footer:
                    for p in s.footer.paragraphs:
                        _sub_placeholders(p)

            # Append remaining sections using template styles
            sorted_sections = sorted(sections, key=lambda s: s.get("order", 0))
            for sec in sorted_sections:
                if sec.get("section_type", "").lower() in replaced_types:
                    continue
                sec_title = sec.get("title", "Section")
                h = doc.add_paragraph()
                if "Heading 1" in doc.styles:
                    h.style = doc.styles["Heading 1"]
                run = h.add_run(sec_title)
                run.bold = True

                self._render_content_block(doc, sec.get("content", ""), normal_style="Normal")

        else:
            # Standard professional academic layout
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_title = p_title.add_run(title)
            run_title.font.size = Pt(20)
            run_title.font.bold = True
            run_title.font.color.rgb = RGBColor(17, 24, 39)

            # Subtitle metadata
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_sub = p_sub.add_run(f"Evidence-Grounded Manuscript Generated by ResearchScribe • Style: {citation_style.upper()}")
            run_sub.font.size = Pt(9.5)
            run_sub.font.italic = True
            run_sub.font.color.rgb = RGBColor(107, 114, 128)

            doc.add_paragraph()  # Spacing

            # Sections in order
            sorted_sections = sorted(sections, key=lambda s: s.get("order", 0))
            for sec in sorted_sections:
                sec_title = sec.get("title", "Section")
                h = doc.add_heading(sec_title, level=1)
                h.paragraph_format.space_before = Pt(14)
                h.paragraph_format.space_after = Pt(6)

                self._render_content_block(doc, sec.get("content", ""))

        # Bibliography section if sources present
        # Bibliography section if sources present
        if sources:
            h_ref = doc.add_heading("References", level=1)
            h_ref.paragraph_format.space_before = Pt(16)
            h_ref.paragraph_format.space_after = Pt(8)

            cited_ids = self._extract_cited_identifiers(sections, claims)
            formatted_refs = citation_formatter.format_cited_bibliography(sources, cited_identifiers=cited_ids, style=citation_style)
            for ref_entry in formatted_refs:
                p_ref = doc.add_paragraph(ref_entry)
                p_ref.paragraph_format.left_indent = Inches(0.4)
                p_ref.paragraph_format.space_after = Pt(4)

        # Optional Verification Appendix
        if include_verification_appendix and claims:
            doc.add_page_break()
            h_app = doc.add_heading("Appendix: Citation & Evidence Provenance Audit", level=1)
            h_app.paragraph_format.space_before = Pt(14)

            p_intro = doc.add_paragraph(
                "This appendix records claim-level evidence verification performed by ResearchScribe. "
                "Every empirical claim is mapped to primary source evidence passages, repository code locations, or verification statuses."
            )
            p_intro.paragraph_format.italic = True

            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Claim Assertion"
            hdr_cells[1].text = "Evidence / Source"
            hdr_cells[2].text = "Status"
            hdr_cells[3].text = "Explanation"

            for cell in hdr_cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True

            for claim in claims:
                row_cells = table.add_row().cells
                row_cells[0].text = claim.get("claim_text", "")
                
                citations = claim.get("citations", [])
                source_str = citations[0].get("in_text_citation", "Unlinked") if citations else "No Citation"
                row_cells[1].text = source_str

                status = claim.get("status", "unverified").upper()
                row_cells[2].text = status
                row_cells[3].text = claim.get("explanation", "") or "Corroborated by repository evidence."

        doc.save(str(out_path))

    def _extract_cited_identifiers(self, sections: List[Dict[str, Any]], claims: Optional[List[Dict[str, Any]]]) -> set:
        """Extracts all source IDs, citation keys, author surnames, and numeric indices cited across the manuscript."""
        cited = set()
        if claims:
            for c in claims:
                for cite in c.get("citations", []):
                    if cite.get("source_id"):
                        cited.add(cite["source_id"])
                    if cite.get("citation_key"):
                        cited.add(cite["citation_key"])
        for sec in sections:
            content = sec.get("content", "")
            for m in re.findall(r'\[(\d+)\]', content):
                cited.add(m)
            for m in re.findall(r'\[([A-Za-z0-9_-]+)\]', content):
                cited.add(m)
            for m in re.findall(r'\(([A-Za-z]+)(?:\s+et\s+al\.)?,?\s+\d{4}\)', content):
                cited.add(m)
            for m in re.findall(r'\((\d+)\)', content):
                cited.add(m)
        return cited

    def _generate_pdf(
        self,
        out_path: Path,
        title: str,
        sections: List[Dict[str, Any]],
        sources: List[Dict[str, Any]],
        claims: Optional[List[Dict[str, Any]]],
        citation_style: str,
        include_verification_appendix: bool,
        abstract: str = "",
        target_journal: str = ""
    ):
        """Generates a publication-grade PDF using reportlab."""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        except ImportError:
            md = self._generate_markdown(title, sections, sources, claims, citation_style, include_verification_appendix)
            out_path.write_text(md, encoding="utf-8")
            return

        doc = SimpleDocTemplate(str(out_path), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            alignment=1,
            spaceAfter=8,
            textColor=colors.HexColor('#111827')
        )
        meta_style = ParagraphStyle(
            'DocMeta',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            alignment=1,
            spaceAfter=18,
            textColor=colors.HexColor('#6B7280')
        )
        h1_style = ParagraphStyle(
            'SectionH1',
            parent=styles['Heading1'],
            fontSize=13,
            leading=17,
            spaceBefore=14,
            spaceAfter=6,
            textColor=colors.HexColor('#1E293B')
        )
        h2_style = ParagraphStyle(
            'SectionH2',
            parent=styles['Heading2'],
            fontSize=11,
            leading=15,
            spaceBefore=10,
            spaceAfter=4,
            textColor=colors.HexColor('#334155')
        )
        body_style = ParagraphStyle(
            'BodyDark',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=14,
            spaceAfter=6,
            textColor=colors.HexColor('#0F172A')
        )

        story.append(Paragraph(title, title_style))
        story.append(Paragraph(f"ResearchScribe Grounded Manuscript &bull; Citation Style: {citation_style.upper()}", meta_style))
        story.append(Spacer(1, 10))

        sorted_sections = sorted(sections, key=lambda s: s.get("order", 0))
        for sec in sorted_sections:
            sec_title = sec.get("title", "Section")
            story.append(Paragraph(sec_title, h1_style))
            content = sec.get("content", "")
            for block in content.split("\n\n"):
                block_clean = block.strip()
                if not block_clean:
                    continue
                if block_clean.startswith("### "):
                    story.append(Paragraph(block_clean.replace("### ", ""), h2_style))
                elif block_clean.startswith("## "):
                    story.append(Paragraph(block_clean.replace("## ", ""), h1_style))
                else:
                    # Clean XML tags for reportlab Paragraph
                    safe_block = block_clean.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(safe_block, body_style))
            story.append(Spacer(1, 6))

        # References
        cited_ids = self._extract_cited_identifiers(sections, claims)
        formatted_refs = citation_formatter.format_cited_bibliography(sources, cited_identifiers=cited_ids, style=citation_style)
        if formatted_refs:
            story.append(Paragraph("References", h1_style))
            for ref in formatted_refs:
                safe_ref = ref.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(safe_ref, body_style))
                story.append(Spacer(1, 3))

        # Verification appendix
        if include_verification_appendix and claims:
            story.append(PageBreak())
            story.append(Paragraph("Appendix: Citation &amp; Evidence Provenance Audit", h1_style))
            table_data = [["Claim Assertion", "Evidence / Source", "Status", "Explanation"]]
            for c in claims:
                c_text = (c.get("claim_text", "")[:100]).replace('&', '&amp;')
                citations = c.get("citations", [])
                src = (citations[0].get("in_text_citation", "Unlinked") if citations else "No Citation").replace('&', '&amp;')
                st = c.get("status", "unverified").upper()
                exp = (c.get("explanation") or "")[:120].replace('&', '&amp;')
                table_data.append([c_text, src, st, exp])

            t = Table(table_data, colWidths=[160, 80, 80, 180])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#0F172A')),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                ('VALIGN', (0,0), (-1,-1), 'TOP')
            ]))
            story.append(t)

        doc.build(story)

    def _render_content_block(self, doc, content: str, normal_style: str = "Normal"):
        """Renders prose, subsection headers, and markdown tables into docx elements."""
        import re
        for block in content.split("\n\n"):
            block_clean = block.strip()
            if not block_clean:
                continue

            # Check if block is a markdown table
            lines = [l.strip() for l in block_clean.splitlines() if l.strip()]
            if len(lines) >= 2 and all(l.startswith("|") and l.endswith("|") for l in lines):
                rows_data = []
                for line in lines:
                    if re.match(r'^\|[\s\-:|]+\|$', line):  # table separator row
                        continue
                    cells = [c.strip() for c in line.strip("|").split("|")]
                    rows_data.append(cells)

                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
                    tbl.style = 'Table Grid'
                    for r_idx, r_data in enumerate(rows_data):
                        for c_idx, val in enumerate(r_data):
                            if c_idx < num_cols:
                                cell = tbl.rows[r_idx].cells[c_idx]
                                cell.text = val
                                if r_idx == 0:
                                    for p in cell.paragraphs:
                                        for run in p.runs:
                                            run.bold = True
                    doc.add_paragraph()
                    continue

            # Regular paragraph or subsections
            if block_clean.startswith("### "):
                h = doc.add_paragraph()
                if "Heading 2" in doc.styles:
                    h.style = doc.styles["Heading 2"]
                run = h.add_run(block_clean.replace("### ", ""))
                run.bold = True
            elif block_clean.startswith("## "):
                h = doc.add_paragraph()
                if "Heading 1" in doc.styles:
                    h.style = doc.styles["Heading 1"]
                run = h.add_run(block_clean.replace("## ", ""))
                run.bold = True
            else:
                p = doc.add_paragraph()
                if normal_style in doc.styles:
                    p.style = doc.styles[normal_style]
                p.add_run(block_clean)

    def _generate_markdown(
        self,
        title: str,
        sections: List[Dict[str, Any]],
        sources: List[Dict[str, Any]],
        claims: Optional[List[Dict[str, Any]]],
        citation_style: str,
        include_verification_appendix: bool
    ) -> str:
        lines = [
            f"# {title}",
            "",
            f"*Generated by ResearchScribe (SRS v3.0 Compliant) • Citation Style: {citation_style.upper()}*",
            "",
            "---",
            ""
        ]

        sorted_sections = sorted(sections, key=lambda s: s.get("order", 0))
        for sec in sorted_sections:
            lines.append(f"## {sec.get('title', 'Section')}")
            lines.append("")
            lines.append(sec.get("content", ""))
            lines.append("")

        if sources:
            lines.append("## References")
            lines.append("")
            cited_ids = self._extract_cited_identifiers(sections, claims)
            formatted_refs = citation_formatter.format_cited_bibliography(sources, cited_identifiers=cited_ids, style=citation_style)
            for ref in formatted_refs:
                lines.append(f"- {ref}")
            lines.append("")

        if include_verification_appendix and claims:
            lines.append("---")
            lines.append("## Appendix: Claim & Evidence Verification Audit")
            lines.append("")
            lines.append("| Claim Assertion | Source Citation | Verification Status | Explanation |")
            lines.append("|---|---|---|---|")
            for c in claims:
                claim_text = c.get("claim_text", "").replace("|", "\\|")
                citations = c.get("citations", [])
                src = citations[0].get("in_text_citation", "Unlinked") if citations else "No Citation"
                status = c.get("status", "unverified")
                exp = (c.get("explanation") or "").replace("|", "\\|")
                lines.append(f"| {claim_text} | {src} | **{status}** | {exp} |")
            lines.append("")

        return "\n".join(lines)

exporter = ManuscriptExporterService()
